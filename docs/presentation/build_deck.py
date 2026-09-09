"""
Build the 90-minute presentation deck (three topics of 30 minutes).

Run from the repository root:

    python docs/presentation/build_deck.py

Slides carry the headline and the key points; the argument to speak lives in the
notes pane, so the deck stays readable on a projector while the detail travels
with the presenter.

Figures are embedded as PNG because PowerPoint's SVG support varies by version.
The PNGs are rendered from the light variants in docs/assets/ by this script
when svglib is available, and reused from docs/presentation/img/ otherwise.
"""

from __future__ import annotations

import glob
import os
from pathlib import Path
from typing import Iterable, List, Optional, Sequence, Tuple

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent.parent
ASSETS = ROOT / "docs" / "assets"
IMG_DIR = Path(__file__).resolve().parent / "img"
OUT = Path(__file__).resolve().parent / "ai-for-smartgrids-apresentacao.pptx"

# 16:9
W, H = Inches(13.333), Inches(7.5)

INK = RGBColor(0x1F, 0x29, 0x33)
MUTED = RGBColor(0x52, 0x60, 0x6D)
FAINT = RGBColor(0x8A, 0x94, 0xA0)
BG = RGBColor(0xFB, 0xFA, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLUE = RGBColor(0x00, 0x72, 0xB2)
GREEN = RGBColor(0x00, 0x9E, 0x73)
AMBER = RGBColor(0xE6, 0x9F, 0x00)
RED = RGBColor(0xD5, 0x5E, 0x00)
DARK = RGBColor(0x06, 0x12, 0x19)

FONT = "Segoe UI"
MONO = "Consolas"


# --------------------------------------------------------------------------
def render_figures() -> None:
    """Rasterise the light SVGs into img/ so the deck can embed them."""
    IMG_DIR.mkdir(parents=True, exist_ok=True)
    try:
        from reportlab.graphics import renderPM
        from svglib.svglib import svg2rlg
    except ImportError:
        print("svglib/reportlab missing - reusing whatever is already in img/")
        return
    for svg in sorted(glob.glob(str(ASSETS / "*-light.svg"))):
        name = os.path.basename(svg).replace("-light.svg", ".png")
        renderPM.drawToFile(svg2rlg(svg), str(IMG_DIR / name), fmt="PNG", dpi=200)
    # The title slide is dark, so it needs the dark banner rather than the light one.
    dark_banner = ASSETS / "banner-dark.svg"
    if dark_banner.exists():
        renderPM.drawToFile(svg2rlg(str(dark_banner)), str(IMG_DIR / "banner-dark.png"),
                            fmt="PNG", dpi=200)


def _tf(shape, text: str, size: int, colour: RGBColor, *, bold=False,
        align=PP_ALIGN.LEFT, font=FONT, spacing=1.0) -> None:
    frame = shape.text_frame
    frame.word_wrap = True
    para = frame.paragraphs[0]
    para.alignment = align
    run = para.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = colour
    run.font.bold = bold
    run.font.name = font
    para.line_spacing = spacing


def textbox(slide, x, y, w, h, text, size, colour, **kwargs):
    box = slide.shapes.add_textbox(x, y, w, h)
    _tf(box, text, size, colour, **kwargs)
    return box


def band(slide, colour: RGBColor, y=Emu(0), h=Inches(0.16)):
    bar = slide.shapes.add_shape(1, Emu(0), y, W, h)  # 1 = rectangle
    bar.fill.solid()
    bar.fill.fore_color.rgb = colour
    bar.line.fill.background()
    bar.shadow.inherit = False
    return bar


def paint(slide, colour: RGBColor = BG):
    bg = slide.shapes.add_shape(1, Emu(0), Emu(0), W, H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = colour
    bg.line.fill.background()
    bg.shadow.inherit = False
    # send to back
    slide.shapes._spTree.remove(bg._element)
    slide.shapes._spTree.insert(2, bg._element)
    return bg


class Deck:
    def __init__(self) -> None:
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self.blank = self.prs.slide_layouts[6]
        self.count = 0

    def _new(self, accent: Optional[RGBColor] = None, bg: RGBColor = BG):
        slide = self.prs.slides.add_slide(self.blank)
        paint(slide, bg)
        if accent is not None:
            band(slide, accent)
        self.count += 1
        return slide

    @staticmethod
    def notes(slide, text: str) -> None:
        slide.notes_slide.notes_text_frame.text = text.strip()

    # -- slide kinds ----------------------------------------------------
    def title_slide(self, image: Path, subtitle: str, author: str, note: str):
        slide = self._new(bg=DARK)
        if image.exists():
            slide.shapes.add_picture(str(image), Emu(0), Inches(0.55), width=W)
        textbox(slide, Inches(0.9), Inches(5.3), Inches(11.5), Inches(0.7),
                subtitle, 22, WHITE, align=PP_ALIGN.CENTER)
        textbox(slide, Inches(0.9), Inches(6.1), Inches(11.5), Inches(0.5),
                author, 15, RGBColor(0x9F, 0xC6, 0xD8), align=PP_ALIGN.CENTER)
        self.notes(slide, note)
        return slide

    def section(self, number: str, title: str, subtitle: str, accent: RGBColor,
                minutes: str, note: str):
        slide = self._new(bg=DARK)
        band(slide, accent, y=Inches(3.05), h=Inches(0.06))
        textbox(slide, Inches(1.0), Inches(2.1), Inches(11.3), Inches(0.9),
                number, 20, accent, bold=True)
        textbox(slide, Inches(1.0), Inches(3.25), Inches(11.3), Inches(1.2),
                title, 44, WHITE, bold=True)
        textbox(slide, Inches(1.0), Inches(4.6), Inches(11.3), Inches(0.8),
                subtitle, 19, RGBColor(0x9F, 0xC6, 0xD8))
        textbox(slide, Inches(1.0), Inches(5.6), Inches(11.3), Inches(0.5),
                minutes, 14, FAINT)
        self.notes(slide, note)
        return slide

    def bullets(self, title: str, items: Sequence, note: str,
                accent: RGBColor = BLUE, subtitle: str = ""):
        slide = self._new(accent)
        textbox(slide, Inches(0.75), Inches(0.5), Inches(11.8), Inches(0.9),
                title, 32, INK, bold=True)
        top = Inches(1.5)
        if subtitle:
            textbox(slide, Inches(0.75), Inches(1.35), Inches(11.8), Inches(0.5),
                    subtitle, 16, MUTED)
            top = Inches(2.0)

        box = slide.shapes.add_textbox(Inches(0.75), top, Inches(11.8),
                                       H - top - Inches(0.6))
        frame = box.text_frame
        frame.word_wrap = True
        # Centre the block vertically: with six or seven points, top alignment
        # leaves the lower half of the slide empty.
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        first = True
        for item in items:
            if isinstance(item, tuple):
                text, level = item
            else:
                text, level = item, 0
            para = frame.paragraphs[0] if first else frame.add_paragraph()
            first = False
            para.level = level
            run = para.add_run()
            is_code = text.startswith("`") and text.endswith("`")
            run.text = ("• " + text) if level == 0 and not is_code else text
            if is_code:
                run.text = text.strip("`")
            run.font.size = Pt(23 - 4 * level if not is_code else 18)
            run.font.name = MONO if is_code else FONT
            run.font.color.rgb = INK if level == 0 else MUTED
            run.font.bold = level == 0 and not is_code
            para.space_after = Pt(16 if level == 0 else 10)
        self.notes(slide, note)
        return slide

    def figure(self, title: str, image: Path, note: str, accent: RGBColor = BLUE,
               caption: str = ""):
        slide = self._new(accent)
        textbox(slide, Inches(0.75), Inches(0.42), Inches(11.8), Inches(0.8),
                title, 28, INK, bold=True)
        if image.exists():
            pic = slide.shapes.add_picture(str(image), Inches(0.6), Inches(1.35),
                                           width=Inches(12.1))
            if pic.height > Inches(5.2):
                scale = Inches(5.2) / pic.height
                pic.height = Inches(5.2)
                pic.width = int(pic.width * scale)
                pic.left = int((W - pic.width) / 2)
        if caption:
            textbox(slide, Inches(0.75), Inches(6.75), Inches(11.8), Inches(0.5),
                    caption, 13, MUTED)
        self.notes(slide, note)
        return slide

    def statement(self, headline: str, detail: str, note: str,
                  accent: RGBColor = RED):
        slide = self._new(bg=DARK)
        band(slide, accent, y=Inches(2.5), h=Inches(0.06))
        textbox(slide, Inches(1.1), Inches(2.85), Inches(11.1), Inches(2.0),
                headline, 36, WHITE, bold=True, spacing=1.15)
        textbox(slide, Inches(1.1), Inches(5.0), Inches(11.1), Inches(1.4),
                detail, 18, RGBColor(0x9F, 0xC6, 0xD8), spacing=1.2)
        self.notes(slide, note)
        return slide

    def table(self, title: str, headers: Sequence[str], rows: Sequence[Sequence[str]],
              note: str, accent: RGBColor = BLUE, widths: Optional[Sequence[float]] = None,
              subtitle: str = "", highlight: Optional[int] = None):
        slide = self._new(accent)
        textbox(slide, Inches(0.75), Inches(0.42), Inches(11.8), Inches(0.8),
                title, 30, INK, bold=True)
        top = Inches(1.45)
        if subtitle:
            textbox(slide, Inches(0.75), Inches(1.3), Inches(11.8), Inches(0.5),
                    subtitle, 15, MUTED)
            top = Inches(1.9)

        n_rows, n_cols = len(rows) + 1, len(headers)
        height = min(Inches(0.46) * n_rows, H - top - Inches(0.5))
        shape = slide.shapes.add_table(n_rows, n_cols, Inches(0.75), top,
                                       Inches(11.8), height)
        tbl = shape.table
        if widths:
            total = sum(widths)
            for i, w in enumerate(widths):
                tbl.columns[i].width = Emu(int(Inches(11.8) * (w / total)))

        for c, head in enumerate(headers):
            cell = tbl.cell(0, c)
            cell.text = head
            para = cell.text_frame.paragraphs[0]
            para.runs[0].font.size = Pt(15)
            para.runs[0].font.bold = True
            para.runs[0].font.color.rgb = WHITE
            para.runs[0].font.name = FONT
            cell.fill.solid()
            cell.fill.fore_color.rgb = accent
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE

        for r, row in enumerate(rows, start=1):
            for c, value in enumerate(row):
                cell = tbl.cell(r, c)
                cell.text = str(value)
                para = cell.text_frame.paragraphs[0]
                para.runs[0].font.size = Pt(14)
                para.runs[0].font.name = FONT
                emphasised = highlight is not None and r - 1 == highlight
                para.runs[0].font.bold = emphasised
                para.runs[0].font.color.rgb = INK if emphasised else MUTED
                cell.fill.solid()
                cell.fill.fore_color.rgb = (
                    RGBColor(0xFD, 0xF1, 0xDC) if emphasised
                    else (WHITE if r % 2 else RGBColor(0xF4, 0xF2, 0xEE))
                )
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        self.notes(slide, note)
        return slide

    def code(self, title: str, lines: Sequence[str], note: str,
             accent: RGBColor = BLUE, caption: str = ""):
        slide = self._new(accent)
        textbox(slide, Inches(0.75), Inches(0.42), Inches(11.8), Inches(0.8),
                title, 30, INK, bold=True)
        panel = slide.shapes.add_shape(1, Inches(0.75), Inches(1.4),
                                       Inches(11.8), Inches(4.7))
        panel.fill.solid()
        panel.fill.fore_color.rgb = DARK
        panel.line.fill.background()
        panel.shadow.inherit = False

        frame = panel.text_frame
        frame.word_wrap = False
        frame.margin_left = Inches(0.35)
        frame.margin_top = Inches(0.28)
        for i, line in enumerate(lines):
            para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            run = para.add_run()
            run.text = line
            run.font.size = Pt(15)
            run.font.name = MONO
            run.font.color.rgb = (
                RGBColor(0x7F, 0xE0, 0xC0) if line.startswith("$")
                else RGBColor(0xE6, 0xED, 0xF3)
            )
            para.line_spacing = 1.15
        if caption:
            textbox(slide, Inches(0.75), Inches(6.35), Inches(11.8), Inches(0.6),
                    caption, 15, MUTED)
        self.notes(slide, note)
        return slide

    def two_up(self, title: str, left_title: str, left: Sequence[str],
               right_title: str, right: Sequence[str], note: str,
               accent: RGBColor = BLUE, left_colour=BLUE, right_colour=AMBER):
        slide = self._new(accent)
        textbox(slide, Inches(0.75), Inches(0.42), Inches(11.8), Inches(0.8),
                title, 30, INK, bold=True)
        for x, head, items, colour in (
            (Inches(0.75), left_title, left, left_colour),
            (Inches(6.95), right_title, right, right_colour),
        ):
            panel = slide.shapes.add_shape(1, x, Inches(1.45), Inches(5.6), Inches(5.0))
            panel.fill.solid()
            panel.fill.fore_color.rgb = WHITE
            panel.line.color.rgb = colour
            panel.line.width = Pt(1.75)
            panel.shadow.inherit = False
            panel.text_frame.text = ""

            textbox(slide, x + Inches(0.3), Inches(1.65), Inches(5.0), Inches(0.5),
                    head, 19, colour, bold=True)
            box = slide.shapes.add_textbox(x + Inches(0.3), Inches(2.25),
                                           Inches(5.0), Inches(4.0))
            frame = box.text_frame
            frame.word_wrap = True
            for i, item in enumerate(items):
                para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
                run = para.add_run()
                run.text = item
                run.font.size = Pt(16)
                run.font.name = FONT
                run.font.color.rgb = MUTED
                para.space_after = Pt(10)
        self.notes(slide, note)
        return slide

    def save(self) -> None:
        self.prs.save(str(OUT))
        print(f"wrote {OUT.relative_to(ROOT)}  ({self.count} slides)")


def main() -> None:
    import sys

    sys.path.insert(0, str(Path(__file__).resolve().parent))
    render_figures()

    from slides import build_slides  # imported late: it imports Deck from here

    deck = Deck()
    build_slides(deck)
    deck.save()


if __name__ == "__main__":
    main()
